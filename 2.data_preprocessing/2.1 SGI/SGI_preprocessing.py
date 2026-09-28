import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DROUGHT_THRESHOLD = -1.0
GAP_SGI_LIMIT = 1.0
MIN_EVENT_MONTHS = 2
MAX_GAP_MONTHS = 2
INPUT_COLUMNS = ["station_name", "date", "Longitude", "Latitude", "sgi"]


def drought_runs(values, months):

    runs = []
    start = None
    for i, value in enumerate(values):
        is_drought = pd.notna(value) and value < DROUGHT_THRESHOLD
        if start is not None and (
            not is_drought or months[i] != months[i - 1] + 1
        ):
            runs.append((start, i - 1))
            start = None
        if is_drought and start is None:
            start = i
    if start is not None:
        runs.append((start, len(values) - 1))
    return runs


def preprocess(data):

    missing = set(INPUT_COLUMNS) - set(data.columns)
    if missing:
        raise ValueError(f"Missing input columns: {sorted(missing)}")
    data = data[INPUT_COLUMNS].copy()
    if data.empty:
        raise ValueError("The input worksheet has no data rows.")
    if data["station_name"].isna().any() or data["date"].isna().any():
        raise ValueError("station_name and date must not be missing.")
    data["date"] = pd.to_datetime(data["date"], errors="raise")
    data["sgi"] = pd.to_numeric(data["sgi"], errors="raise")
    if data["sgi"].isin([float("inf"), float("-inf")]).any():
        raise ValueError("SGI must be finite or missing.")
    data = data.sort_values(["station_name", "date"]).reset_index(drop=True)
    data["_month"] = data["date"].dt.to_period("M").astype("int64")
    if data.duplicated(["station_name", "_month"]).any():
        raise ValueError("Expected at most one observation per station/month.")

    output = data.rename(columns={"sgi": "sgi_original_file"})
    output["sgi"] = output["sgi_original_file"]
    output["is_drought"] = False
    output["event_id"] = pd.Series(pd.NA, index=output.index, dtype="string")
    output["adjustment_type"] = "unchanged"
    summary = {"preliminary_events": 0, "merged_pairs": 0,
               "adjusted_months": 0, "final_events": 0}

    for _, group in output.groupby("station_name", sort=False):
        positions = group.index.to_list()
        values = group["sgi_original_file"].to_list()
        months = group["_month"].to_list()
        runs = drought_runs(values, months)
        summary["preliminary_events"] += len(runs)


        for (left_start, left_end), (right_start, right_end) in zip(runs, runs[1:]):
            gap = right_start - left_end - 1
            if (
                left_end - left_start + 1 < MIN_EVENT_MONTHS
                or right_end - right_start + 1 < MIN_EVENT_MONTHS
                or not 1 <= gap <= MAX_GAP_MONTHS
                or months[right_start] - months[left_end] != gap + 1
            ):
                continue
            if not all(pd.notna(v) and v < GAP_SGI_LIMIT
                       for v in values[left_end + 1:right_start]):
                continue

            first_gap = positions[left_end + 1]
            output.at[first_gap, "sgi"] = values[left_end]
            output.at[first_gap, "adjustment_type"] = (
                "1-month gap: copied previous drought SGI" if gap == 1 else
                "2-month gap month 1: copied previous drought SGI"
            )
            if gap == 2:
                second_gap = positions[left_end + 2]
                output.at[second_gap, "sgi"] = values[right_start]
                output.at[second_gap, "adjustment_type"] = (
                    "2-month gap month 2: copied next drought SGI"
                )
            summary["merged_pairs"] += 1
            summary["adjusted_months"] += gap

        adjusted = output.loc[positions, "sgi"].to_list()
        for start, end in drought_runs(adjusted, months):
            summary["final_events"] += 1
            event_rows = positions[start:end + 1]
            output.loc[event_rows, "is_drought"] = True
            output.loc[event_rows, "event_id"] = f"Event_{summary['final_events']}"

    return output.drop(columns="_month"), summary


def save_excel(data, path):

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        data.to_excel(writer, sheet_name="Location_SGI_Final", index=False)
        sheet = writer.sheets["Location_SGI_Final"]
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30
        widths = [26, 14, 16, 16, 21, 13, 16, 16, 58]
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
            header = sheet.cell(1, column)
            header.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            header.fill = PatternFill("solid", fgColor="24465C")
            header.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            row[1].number_format = "yyyy-mm-dd"
            for cell in row[2:4]:
                cell.number_format = "0.000000"
            for cell in row[4:6]:
                cell.number_format = "0.000"


def main():
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", type=Path,
                        default=script_dir.parents[1] / "1.Input data" / "Location_SGI.xlsx")
    parser.add_argument("--output", type=Path,
                        default=script_dir / "Location_SGI_final.xlsx")
    args = parser.parse_args()
    if args.input.resolve() == args.output.resolve():
        parser.error("Input and output must be different files.")
    data = pd.read_excel(args.input, sheet_name="Location_SGI", engine="openpyxl")
    output, summary = preprocess(data)
    save_excel(output, args.output)
    print(f"Saved: {args.output.resolve()}")
    print(f"Stations: {output['station_name'].nunique():,}; records: {len(output):,}")
    for label, count in summary.items():
        print(f"{label}: {count:,}")


if __name__ == "__main__":
    main()
